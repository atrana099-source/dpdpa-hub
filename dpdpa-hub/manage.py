"""Command line helper.  Usage:
  python manage.py init            create tables, admin user, default settings and a starter notice
  python manage.py seed-tracker    import the Carzonrent gap/ROPA/retention tracker from seed/tracker.xlsx
  python manage.py demo            add sample consents, a request and an incident (for trials only)
  python manage.py create-user USERNAME ROLE   (password asked or from NEW_USER_PASSWORD)
"""
import datetime
import os
import re
import secrets
import sys

from app import create_app, db
from app.auth import create_user
from app.util import now_iso, set_setting, get_setting

HERE = os.path.dirname(os.path.abspath(__file__))

PURPOSES = [
    ("BOOKING", "Booking and trip delivery", "Use your details to confirm and deliver your chauffeur-driven booking.", "Name, mobile, email, pickup and drop locations, travel dates", "Consent (s.6)", "Trip duration plus the legal minimum"),
    ("TRIP_COMM", "Trip updates", "Send chauffeur details, ETA and trip changes by SMS, WhatsApp or email.", "Name, mobile, email, trip details", "Consent (s.6)", "Trip duration plus 90 days"),
    ("BILLING", "Billing and invoicing", "Raise invoices and keep tax records for your booking.", "Name, GSTIN, billing address, amounts", "Consent (s.6)", "As required by tax law"),
    ("KYC", "Identity check for security", "Check your identity when required by your company or for safety.", "ID type and number, photo", "Consent (s.6)", "Trip duration plus 30 days"),
    ("FEEDBACK", "Feedback and service quality", "Ask for feedback and use it to improve service.", "Name, contact, trip details, rating", "Consent (s.6)", "24 months"),
    ("MARKETING", "Offers and promotions", "Send offers and news about our services.", "Name, email, mobile", "Consent (s.6)", "Until you withdraw"),
]
DEPT_MAP = {"Department": "name", "Interview date": "interview_date", "Preliminary rating": "rating", "HOD declaration signed": "declaration_signed",
            "Evidence pack received": "evidence_received", "Key issue": "key_issue"}
MAPS = {
    "gaps": ("Gap Register", {"ID": "code", "Department": "department", "Theme": "theme", "Finding": "finding", "Type": "type", "Priority": "priority",
                              "DPDP reference": "dpdp_ref", "Source": "source", "Action ID": "action_code", "Status": "status",
                              "Target date": "target_date", "Named owner": "owner", "Evidence / remarks": "remarks"}),
    "actions": ("Action Tracker", {"ID": "code", "Action": "action", "Owner (function)": "owner_function", "Phase": "phase", "Due date": "due_date",
                                   "Departments in scope": "departments", "Status": "status", "% complete": "pct", "Named owner": "named_owner",
                                   "Evidence / link": "evidence", "Remarks": "remarks"}),
    "retention": ("Retention Schedule", {"ID": "code", "Data category": "data_category", "Departments": "departments", "Data principals": "data_principals",
                                         "Clock starts": "clock_starts", "Statutory minimum": "statutory_min", "Proposed retention": "proposed",
                                         "Basis type": "basis_type", "Disposal": "disposal", "Current practice (interviews)": "current_practice",
                                         "ROPA ID": "ropa_ids", "Owner": "owner", "Status": "status"}),
    "ropa": ("ROPA", {"ROPA ID": "code", "Department": "department", "Processing activity": "activity", "Data principals": "data_principals",
                      "Data categories": "data_categories", "Purpose": "purpose", "Proposed lawful basis": "lawful_basis", "Systems": "systems",
                      "Recipients": "recipients", "Cross-border": "cross_border", "Retention (as stated)": "retention",
                      "Retention gap?": "retention_gap", "Owner (HOD)": "owner", "Last reviewed": "last_reviewed"}),
    "departments": ("Departments", DEPT_MAP),
}
DATE_COLS = {"target_date", "due_date", "interview_date", "last_reviewed", "next_review"}


def clean(col, v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if col == "pct":
        try:
            f = float(v)
            return int(round(f * 100 if f <= 1 else f))
        except (TypeError, ValueError):
            return 0
    s = str(v).strip()
    if col in DATE_COLS and s and not re.match(r"^\d{4}-\d{2}-\d{2}", s):
        for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d-%b-%Y"):
            try:
                return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
        return ""
    return s


def cmd_init():
    db.init_schema()
    defaults = {"org_name": os.environ.get("ORG_NAME", "Carzonrent India Pvt Ltd"), "dpo_name": os.environ.get("DPO_NAME", "Data Protection Officer"),
                "dpo_email": os.environ.get("DPO_EMAIL", "dpo@example.com"), "dsr_sla_days": "30"}
    for k, v in defaults.items():
        if db.scalar("SELECT COUNT(*) FROM settings WHERE key = ?", (k,)) == 0:
            set_setting(k, v)
    if db.scalar("SELECT COUNT(*) FROM users") == 0:
        user = os.environ.get("ADMIN_USER", "admin")
        pw = os.environ.get("ADMIN_PASSWORD")
        generated = not pw
        pw = pw or secrets.token_urlsafe(12)
        create_user(user, pw, "admin", "Administrator")
        print(f"Admin user created: {user}")
        if generated:
            print(f"Generated password (shown once, change it after login): {pw}")
    if db.scalar("SELECT COUNT(*) FROM purposes") == 0:
        for code, name, desc, cats, basis, ret in PURPOSES:
            db.execute("INSERT INTO purposes (code, name, description, data_categories, legal_basis, retention, status) VALUES (?,?,?,?,?,?, 'Active')",
                       (code, name, desc, cats, basis, ret))
    if db.scalar("SELECT COUNT(*) FROM notices") == 0:
        from app.views_consent import DEFAULT_BODY
        from app import consent_service as cs
        codes = ["BOOKING", "TRIP_COMM", "BILLING", "FEEDBACK", "MARKETING"]
        n = {"slug": "guest-booking", "title": "How we use your personal data for your booking", "body": DEFAULT_BODY}
        purposes = [db.one("SELECT * FROM purposes WHERE code = ?", (c,)) for c in codes]
        db.insert("INSERT INTO notices (slug, version, title, language, body, purpose_codes, status, content_hash, created_at, published_at) "
                  "VALUES (?,?,?,?,?,?, 'Published', ?, ?, ?)",
                  (n["slug"], 1, n["title"], "English", n["body"], ",".join(codes), cs.notice_hash(n["title"], n["body"], purposes), now_iso(), now_iso()))
        print("Starter notice 'guest-booking' published. Review the text with Legal.")
    print("Initialised.")


def cmd_seed_tracker(path=None):
    import openpyxl
    path = path or os.path.join(HERE, "seed", "tracker.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    for table, (sheet, mapping) in MAPS.items():
        if sheet not in wb.sheetnames:
            print(f"skip {sheet}: not in workbook")
            continue
        ws = wb[sheet]
        header_row = next((r for r in range(1, 12) if sum(1 for c in ws[r] if c.value in mapping) >= 3), None)
        if not header_row:
            print(f"skip {sheet}: header row not found")
            continue
        cols = {c.column - 1: mapping[c.value] for c in ws[header_row] if c.value in mapping}
        key = "name" if table == "departments" else "code"
        n = 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            rec = {col: clean(col, row[i]) for i, col in cols.items()}
            if not rec.get(key):
                continue
            if table == "departments" and (rec[key].lower() == "total" or not rec.get("rating")):
                continue
            if table != "departments" and not re.match(r"^[A-Za-z]{1,3}-?\d+$", rec[key]):
                continue
            exists = db.scalar(f"SELECT COUNT(*) FROM {table} WHERE {key} = ?", (rec[key],))
            names = list(rec)
            if exists:
                db.execute(f"UPDATE {table} SET " + ", ".join(f"{c}=?" for c in names if c != key) + f" WHERE {key} = ?",
                           [rec[c] for c in names if c != key] + [rec[key]])
            else:
                db.insert(f"INSERT INTO {table} (" + ",".join(names) + ") VALUES (" + ",".join("?" * len(names)) + ")", [rec[c] for c in names])
            n += 1
        print(f"{table}: {n} rows")


def cmd_demo():
    from app import consent_service as cs
    from app.views_dsr import create_request
    cs.record_consent("demo.guest@example.com", "guest-booking", ["BOOKING", "TRIP_COMM"], "English", "demo")
    cs.record_consent("9876543210", "guest-booking", ["BOOKING", "MARKETING"], "English", "demo")
    cs.withdraw_consent("9876543210", "guest-booking", ["MARKETING"], "demo")
    create_request("Erasure", "Demo Guest", "demo.guest@example.com", "demo.guest@example.com", "Please erase my data.")
    aware = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(hours=20)).strftime("%Y-%m-%dT%H:%M:%SZ")
    from app.util import next_ref
    db.insert("INSERT INTO incidents (ref, title, description, aware_at, severity, status, created_at) VALUES (?,?,?,?, 'High', 'Investigating', ?)",
              (next_ref("INC", "incidents"), "Demo: itinerary emailed to wrong recipient", "Sample record for training.", aware, now_iso()))
    print("Demo data added.")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    app = create_app()
    cmd = sys.argv[1]
    with app.app_context():
        if cmd == "init":
            cmd_init()
        elif cmd == "seed-tracker":
            cmd_seed_tracker(sys.argv[2] if len(sys.argv) > 2 else None)
        elif cmd == "demo":
            cmd_demo()
        elif cmd == "create-user" and len(sys.argv) >= 4:
            pw = os.environ.get("NEW_USER_PASSWORD") or input("Password: ")
            create_user(sys.argv[2], pw, sys.argv[3])
            print("User created.")
        else:
            print(__doc__)
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
